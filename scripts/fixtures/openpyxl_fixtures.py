"""Generates the openpyxl fixtures (run by scripts/make-fixtures.sh)."""
import datetime, sys
import openpyxl
from openpyxl.styles import Font
from openpyxl.utils.datetime import CALENDAR_MAC_1904

out = sys.argv[1]

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "People"
ws.append(["name", "age", "joined", "active", "score"])
ws.append(["Alice", 30, datetime.datetime(2020, 5, 17, 8, 15), True, 9.5])
ws.append(["Bob", None, datetime.date(1999, 12, 31), False, -1])
ws.append(["Čeněk & <Co>", 41, datetime.datetime(1900, 1, 1), True, 1e-7])
ws.append(["_x0041_ literal", 0, None, None, 12345678901])
for c in ws[1]:
    c.font = Font(bold=True)
ws.merge_cells("A7:B7")
ws["A7"] = "merged"
other = wb.create_sheet("Other")
other.append(["k", "v"])
other.append(["x", 1])
wb.save(f"{out}/openpyxl.xlsx")

wb = openpyxl.Workbook()
wb.epoch = CALENDAR_MAC_1904
ws = wb.active
ws.append(["when"])
ws.append([datetime.datetime(2026, 10, 9, 6, 0)])
ws.append([datetime.date(1904, 1, 2)])
wb.save(f"{out}/openpyxl-1904.xlsx")
